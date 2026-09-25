
from io import BytesIO
from html import escape
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
from reportlab.lib.enums import TA_CENTER
import os

def _safe(x):
    return "" if x is None else str(x)

def _sections(case,result):
    econ=result["economics"]; bia=result["budget_impact"]
    return [
        ("Executive Summary", [
            f"Evidence readiness: {result['readiness_score']:.1f}/100",
            f"Overall uncertainty: {result['uncertainty_label']}",
            result["interpretation"]
        ]),
        ("Technology Identification", [
            f"Device: {_safe(case.get('identity',{}).get('name'))}",
            f"Manufacturer: {_safe(case.get('identity',{}).get('manufacturer'))}",
            f"Model/version: {_safe(case.get('identity',{}).get('model'))}",
            f"Type/Class: {_safe(case.get('identity',{}).get('device_type'))} / {_safe(case.get('identity',{}).get('class'))}",
            f"Intended purpose: {_safe(case.get('identity',{}).get('intended_purpose'))}",
            f"Indication: {_safe(case.get('identity',{}).get('indication'))}",
        ]),
        ("JCA / Eligibility", [
            f"JCA available: {'Yes' if case.get('eligibility',{}).get('jca_available') else 'No'}",
            f"JCA reference: {_safe(case.get('eligibility',{}).get('jca_ref'))}",
            f"National pathway: {_safe(case.get('eligibility',{}).get('national_pathway'))}",
        ]),
        ("PICO", [
            f"Population: {_safe(case.get('pico',{}).get('population'))}",
            f"Intervention: {_safe(case.get('pico',{}).get('intervention'))}",
            f"Comparator: {_safe(case.get('pico',{}).get('comparator'))}",
            f"Outcomes: {_safe(case.get('pico',{}).get('outcomes'))}",
        ]),
        ("Clinical Assessment", [
            f"Study design: {_safe(case.get('clinical',{}).get('study_design'))}",
            f"Certainty: {_safe(case.get('clinical',{}).get('certainty'))}",
            f"Rationale: {_safe(case.get('clinical',{}).get('notes'))}",
            f"Clinical safety rationale: {_safe(case.get('safety',{}).get('notes'))}",
        ]),
        ("Economic Evaluation", [
            f"Method: {_safe(case.get('economic',{}).get('analysis_type'))}",
            f"Upfront fixed cost: {econ['upfront_fixed']:.2f} €",
            f"Annualized fixed cost: {econ['annualized_fixed']:.2f} €",
            f"Device cost per case: {econ['device_cost_per_case']:.2f} €",
            f"Incremental cost per case: {econ['incremental_cost_per_case']:.2f} €",
            f"ICER: {'Not defined' if econ['icer'] is None else f'{econ['icer']:.2f} € per unit of effect'}",
        ]),
        ("Budget Impact", [
            f"Year 1 treated: {bia['y1_users']:.1f}",
            f"Year 1 budget impact: {bia['year1']:.2f} €",
            f"Year 3 treated: {bia['y3_users']:.1f}",
            f"Year 3 budget impact: {bia['year3']:.2f} €",
        ]),
        ("Domain Profile", [f"{k}: {v:.1f}/100" for k,v in result["domain_scores"].items()]),
        ("Red Flags", result["red_flags"] or ["None identified from current inputs."]),
        ("Evidence Gaps", result["gaps"]),
        ("Source Traceability", [f"{k}: {v}" for k,v in case.get("sources",{}).items() if v] or ["No sources entered."]),
    ]

def build_docx_bytes(case,result):
    doc=Document()
    styles=doc.styles
    styles["Normal"].font.name="Arial"
    styles["Normal"].font.size=Pt(10.5)
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("MedTech HTA Assessment Report")
    r.bold=True; r.font.size=Pt(18)
    p2=doc.add_paragraph()
    p2.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run("Greek MedTech HTA Workspace").italic=True
    for title,items in _sections(case,result):
        doc.add_heading(title,level=1)
        for item in items:
            if not item: continue
            p=doc.add_paragraph(style="List Bullet")
            p.add_run(_safe(item))
    doc.add_heading("Methodological Note",level=1)
    doc.add_paragraph("This report is a decision-support output and does not constitute an official regulatory approval, reimbursement decision or HTA authority decision.")
    bio=BytesIO(); doc.save(bio); return bio.getvalue()

def build_pdf_bytes(case,result):
    bio=BytesIO()
    styles=getSampleStyleSheet()
    # Unicode font for Greek/English reports.
    font_path="/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    bold_path="/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
    font_name="Helvetica"
    bold_name="Helvetica-Bold"
    if os.path.exists(font_path):
        try:
            pdfmetrics.registerFont(TTFont("HTAUnicode",font_path))
            font_name="HTAUnicode"
            if os.path.exists(bold_path):
                pdfmetrics.registerFont(TTFont("HTAUnicodeBold",bold_path))
                bold_name="HTAUnicodeBold"
        except Exception:
            pass
    title=ParagraphStyle("T",parent=styles["Title"],alignment=TA_CENTER,fontSize=16,leading=20,fontName=bold_name)
    h=ParagraphStyle("H",parent=styles["Heading2"],fontName=bold_name)
    body=ParagraphStyle("B",parent=styles["BodyText"],fontName=font_name,leading=14)
    story=[Paragraph("MedTech HTA Assessment Report",title),Spacer(1,12)]
    for sec,items in _sections(case,result):
        story.append(Paragraph(escape(sec),h))
        for item in items:
            story.append(Paragraph("• "+escape(_safe(item)),body))
            story.append(Spacer(1,4))
        story.append(Spacer(1,8))
    story.append(Paragraph("Methodological Note",h))
    story.append(Paragraph("Decision-support output only; not an official regulatory or reimbursement decision.",body))
    doc=SimpleDocTemplate(bio,pagesize=A4,rightMargin=40,leftMargin=40,topMargin=45,bottomMargin=45)
    doc.build(story)
    return bio.getvalue()
